"""Read Office source cells; formulas are preserved but never evaluated."""
import io
import posixpath
import re
import zipfile
from pathlib import PurePath
from xml.etree import ElementTree as ET

from .domain import Cell, Column, ParsedTable, validate_tables

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def unit_from_header(value):
    match = re.search(r'[（(](元|万元|亿元|%|％|个|人|天|小时|kg|g|m|cm)[）)]\s*$', value)
    return match.group(1) if match else None


def from_rows(table_id, rows, *, name='', locator=None, unknown=False, row_numbers=None):
    width = max((len(row) for row in rows), default=0)
    headers = rows[0] if rows else []
    columns = tuple(Column(column_id=f'c-{i + 1}', index=i + 1,
                           header_path=(str(headers[i]),) if i < len(headers) and headers[i] is not None else ()) for i in range(width))
    numbers = row_numbers if row_numbers is not None else list(range(1, len(rows)+1))
    cells = tuple(Cell(row=r, column=c, raw_text='' if value is None else str(value),
                       value_kind='empty' if value is None else 'unknown' if unknown else 'text',
                       value=None if value is None else str(value),
                       unit=unit_from_header(str(headers[c-1])) if r > 1 and c <= len(headers) else None)
                  for r, row in zip(numbers, rows, strict=True) for c, value in enumerate(row, 1))
    return ParsedTable(table_id=table_id, name=name, row_count=max(numbers, default=0), columns=columns, cells=cells, locator=locator or {})


def docx_table(table, table_id, locator):
    rows = [[c.text for c in r.cells] for r in table.rows]
    base = from_rows(table_id, rows, locator=locator)
    owners = {}
    for r, row in enumerate(table.rows, 1):
        for c, cell in enumerate(row.cells, 1):
            owners.setdefault(cell._tc, []).append((r, c))
    by_position = {(c.row, c.column): c for c in base.cells}
    cells = []
    for positions in owners.values():
        top, left = min(positions)
        bottom, right = max(r for r, _ in positions), max(c for _, c in positions)
        if len(positions) != (bottom - top + 1) * (right - left + 1):
            raise ValueError('nonrectangular merged source cell')
        cells.append(by_position[top, left].model_copy(update={'row_span': bottom-top+1, 'column_span': right-left+1}))
    return ParsedTable.model_validate({**base.model_dump(), 'cells': cells})


def pptx_table(table, table_id, locator):
    base = from_rows(table_id, [[c.text for c in r.cells] for r in table.rows], locator=locator)
    cells = []
    for cell in base.cells:
        original = table.cell(cell.row-1, cell.column-1)
        if original.is_spanned:
            continue
        cells.append(cell.model_copy(update={'row_span': original.span_height, 'column_span': original.span_width}))
    return ParsedTable.model_validate({**base.model_dump(), 'cells': cells})


def xlsx_tables(data):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
    from openpyxl.styles.numbers import is_date_format, BUILTIN_FORMATS
    from openpyxl.utils.datetime import from_excel
    # The library supplies date formats, not numeric precision; raw <v> stays a string.
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    tables = []
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            workbook = ET.fromstring(archive.read('xl/workbook.xml'))
            styles = ET.fromstring(archive.read('xl/styles.xml'))
            formats = {int(e.attrib['numFmtId']): e.attrib['formatCode'] for e in styles.iter(NS + 'numFmt')}
            style_formats = [formats.get(int(e.attrib.get('numFmtId', 0)),
                                          BUILTIN_FORMATS.get(int(e.attrib.get('numFmtId', 0)), 'General'))
                             for e in styles.find(NS + 'cellXfs')]
            rels = {e.attrib['Id']: e.attrib['Target'] for e in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
            shared = []
            if 'xl/sharedStrings.xml' in archive.namelist():
                shared = [''.join(item.itertext()) for item in ET.fromstring(archive.read('xl/sharedStrings.xml'))]
            for index, sheet in enumerate(workbook.find(NS + 'sheets'), 1):
                target = rels[sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
                path = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
                if not path.startswith('xl/') or '..' in path.split('/'):
                    raise ValueError('invalid worksheet path')
                root = ET.fromstring(archive.read(path))
                if root.tag != NS + 'worksheet':
                    continue  # Chartsheets are not cell tables.
                if any(e.tag == NS + 'dimension' and ':' in e.attrib.get('ref', '') and
                       (range_boundaries(e.attrib['ref'])[2] * range_boundaries(e.attrib['ref'])[3] > 2000000)
                       for e in root):
                    raise ValueError('worksheet extent limit')
                merges = [range_boundaries(e.attrib['ref']) for e in root.iter(NS + 'mergeCell')]
                cells, height, width = [], 0, 0
                for element in root.iter(NS + 'c'):
                    row, col = coordinate_to_tuple(element.attrib['r'])
                    height, width = max(height, row), max(width, col)
                    merge = next((m for m in merges if m[0] <= col <= m[2] and m[1] <= row <= m[3]), None)
                    if merge and (col, row) != merge[:2]:
                        continue
                    raw = element.find(NS + 'v')
                    raw_value = raw.text if raw is not None else None
                    typ = element.attrib.get('t', 'n')
                    kind, value = 'empty', None
                    if typ == 'inlineStr':
                        kind, value = 'text', ''.join(e.text or '' for e in element.iter(NS + 't'))
                    elif typ == 's' and raw_value is not None:
                        kind, value = 'text', shared[int(raw_value)]
                    elif typ in ('str', 'd') and raw is not None:
                        kind, value = ('date' if typ == 'd' else 'text'), raw_value or ''
                    elif raw_value is not None:
                        kind, value = {'b': 'boolean', 'e': 'error'}.get(typ, 'number'), raw_value
                        if kind == 'boolean':
                            value = 'true' if raw_value == '1' else 'false'
                        style = int(element.attrib.get('s', '0'))
                        if style < 0 or style >= len(style_formats):
                            raise ValueError('invalid cell style')
                        if kind == 'number' and is_date_format(style_formats[style]):
                            kind, value = 'date', from_excel(float(raw_value), book.epoch).isoformat()
                    formula = element.find(NS + 'f')
                    state = 'not_formula' if formula is None else 'cache_present' if raw_value is not None else 'cache_missing'
                    if formula is not None and not formula.text:
                        state = 'unknown'  # Shared/array formula with no local expression.
                    cells.append(Cell(row=row, column=col, raw_text=value if kind == 'text' else raw_value or '',
                        value_kind=kind, value=value, formula=None if formula is None else '=' + (formula.text or ''), formula_state=state,
                        row_span=merge[3]-row+1 if merge else 1, column_span=merge[2]-col+1 if merge else 1))
                    if len(cells) > 200000:
                        raise ValueError('table cell limit')
                for left, top, right, bottom in merges:
                    height, width = max(height, bottom), max(width, right)
                headers = {c.column: c.raw_text for c in cells if c.row == 1}
                cells = [c.model_copy(update={'unit': unit_from_header(headers.get(c.column, ''))}) if c.row > 1 else c for c in cells]
                tables.append(ParsedTable(table_id=f't-{index}', name=sheet.attrib['name'], row_count=height,
                    columns=tuple(Column(column_id=f'c-{c}', index=c, header_path=(headers[c],) if c in headers else ()) for c in range(1, width+1)),
                    cells=tuple(cells), locator={'sheet': sheet.attrib['name']}))
    finally:
        book.close()
    return validate_tables(tables)


def extract_tables(content):
    ext = PurePath(content.filename).suffix.lower()
    if ext == '.xlsx':
        return xlsx_tables(content.data)
    if ext == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(content.data))
        return validate_tables(docx_table(t, f't-{i}', {'table': i})
                               for i, t in enumerate(doc.tables, 1))
    if ext == '.pptx':
        from pptx import Presentation
        tables = []
        for page, slide in enumerate(Presentation(io.BytesIO(content.data)).slides, 1):
            def walk(shapes):
                for shape in shapes:
                    if hasattr(shape, 'shapes'):
                        walk(shape.shapes)
                    elif shape.has_table:
                        tables.append(pptx_table(shape.table, f't-{len(tables)+1}', {'slide': page, 'shape': shape.shape_id}))
            walk(slide.shapes)
        return validate_tables(tables)
    return ()
