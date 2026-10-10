"""Bounded structural extraction; no macros, external links, OCR or model calls."""

from __future__ import annotations

import io
import json
import re
import zipfile
from pathlib import PurePath

from .domain import DocumentParser, ParsedDocument, Section, SourceContent, UnsupportedDocument

FEISHU_NATIVE_TYPE = "application/vnd.suite.feishu+json"


def preflight(content: SourceContent, *, max_bytes=30 * 1024 * 1024):
    """Shared format checks before any enhanced engine sees the bytes."""
    if not content.data or len(content.data) > min(max_bytes, 30 * 1024 * 1024):
        raise UnsupportedDocument("parser_input_limit")
    extension = PurePath(content.filename).suffix.lower()
    types = {'.pdf': 'application/pdf', '.png': 'image/png', '.jpg': 'image/jpeg',
             '.jpeg': 'image/jpeg', '.tif': 'image/tiff', '.tiff': 'image/tiff',
             '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
             '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
             '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation'}
    if extension in types and content.media_type not in {types[extension], 'application/octet-stream'}:
        raise UnsupportedDocument("parser_format_mismatch")
    signatures = {'.pdf': (b'%PDF-',), '.png': (b'\x89PNG\r\n\x1a\n',), '.jpg': (b'\xff\xd8\xff',),
                  '.jpeg': (b'\xff\xd8\xff',), '.tif': (b'II*\x00', b'MM\x00*'), '.tiff': (b'II*\x00', b'MM\x00*')}
    if extension in signatures and not any(content.data.startswith(signature) for signature in signatures[extension]):
        raise UnsupportedDocument("parser_format_mismatch")
    if extension in {'.docx', '.xlsx', '.pptx'}:
        try:
            StructuredDocumentParser()._check_zip(content.data)
            with zipfile.ZipFile(io.BytesIO(content.data)) as archive:
                prefix = {'.docx': 'word/', '.xlsx': 'xl/', '.pptx': 'ppt/'}[extension]
                if '[Content_Types].xml' not in archive.namelist() or not any(name.startswith(prefix) for name in archive.namelist()):
                    raise UnsupportedDocument("parser_format_mismatch")
        except (zipfile.BadZipFile, OSError):
            raise UnsupportedDocument("parser_invalid_office") from None
    if extension == '.pdf':
        from pypdf import PdfReader
        try:
            if PdfReader(io.BytesIO(content.data)).is_encrypted:
                raise UnsupportedDocument("parser_encrypted_pdf")
        except UnsupportedDocument:
            raise
        except Exception:
            raise UnsupportedDocument("parser_invalid_pdf") from None


class StructuredDocumentParser(DocumentParser):
    def __init__(self, *, max_bytes: int = 30 * 1024 * 1024,
                 max_expanded_bytes: int = 150 * 1024 * 1024, max_sections: int = 50000):
        self.max_bytes = max_bytes
        self.max_expanded_bytes = max_expanded_bytes
        self.max_sections = max_sections

    def parse(self, content: SourceContent) -> ParsedDocument:
        if len(content.data) > self.max_bytes:
            raise UnsupportedDocument("文件超过解析大小限制")
        ext = PurePath(content.filename).suffix.lower()
        try:
            if content.media_type == FEISHU_NATIVE_TYPE:
                result = self._native(json.loads(content.data))
            elif ext in {".md", ".txt"}:
                result = self._text(content.data, markdown=ext == ".md")
            elif ext in {".docx", ".xlsx", ".pptx"}:
                self._check_zip(content.data)
                result = getattr(self, "_" + ext[1:])(content.data)
            elif ext == ".pdf":
                result = self._pdf(content.data)
            else:
                raise UnsupportedDocument(f"暂不支持此格式：{ext or content.media_type}")
        except UnsupportedDocument:
            raise
        except ImportError as exc:
            raise UnsupportedDocument(f"缺少格式解析依赖：{exc.name}") from exc
        except Exception as exc:
            raise UnsupportedDocument(f"文档结构无效或无法解析：{type(exc).__name__}") from exc
        if not result.sections:
            detail = "；".join(result.warnings)
            raise UnsupportedDocument("未发现可检索文字，图片/扫描内容尚未实现 OCR" + (f"；{detail}" if detail else ""))
        if len(result.sections) > self.max_sections:
            raise UnsupportedDocument("文档段落数量超过解析限制")
        if sum(len(section.text) for section in result.sections) > self.max_expanded_bytes:
            raise UnsupportedDocument("解析文本超过限制")
        return result

    def _check_zip(self, data: bytes) -> None:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 20000 or sum(item.file_size for item in infos) > self.max_expanded_bytes:
                raise UnsupportedDocument("Office 压缩包展开后超过限制")
            if any(item.flag_bits & 1 for item in infos):
                raise UnsupportedDocument("不支持加密文档")
            if any(item.filename.lower().endswith("vbaproject.bin") for item in infos):
                raise UnsupportedDocument("不接收含宏的 Office 文件")

    def _text(self, data: bytes, *, markdown: bool) -> ParsedDocument:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise UnsupportedDocument("文本文件需要 UTF-8 编码") from exc
        if "\x00" in text:
            raise UnsupportedDocument("内容不是有效文本")
        sections, paragraph, start, heading = [], [], 1, ""
        lines = text.splitlines()
        def flush(end: int) -> None:
            if paragraph:
                sections.append(Section("\n".join(paragraph), {"line_start": start, "line_end": end}, heading))
                paragraph.clear()
        for number, line in enumerate(lines, 1):
            if markdown and re.match(r"^#{1,6}\s+", line):
                flush(number - 1)
                heading = re.sub(r"^#{1,6}\s+", "", line).strip()
            if not line.strip():
                flush(number - 1)
            else:
                if not paragraph:
                    start = number
                paragraph.append(line)
        flush(len(lines))
        warnings = ("Markdown 图片仅保留引用文本，未执行 OCR",) if markdown and re.search(r"!\[", text) else ()
        return ParsedDocument(tuple(sections), warnings)

    def _docx(self, data: bytes) -> ParsedDocument:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(io.BytesIO(data))
        sections, warnings, heading, paragraph_number, table_number = [], [], "", 0, 0
        for element in doc.element.body.iterchildren():
            tag = element.tag.rsplit("}", 1)[-1]
            if tag == "p":
                paragraph = Paragraph(element, doc)
                paragraph_number += 1
                text = paragraph.text.strip()
                if paragraph.style and paragraph.style.name.startswith("Heading"):
                    heading = text
                if text:
                    sections.append(Section(text, {"paragraph": paragraph_number}, heading))
            elif tag == "tbl":
                table_number += 1
                table = Table(element, doc)
                headers = [cell.text.strip() for cell in table.rows[0].cells] if table.rows else []
                for row_number, row in enumerate(table.rows, 1):
                    values = [cell.text.strip() for cell in row.cells]
                    if any(values):
                        sections.append(Section(self._row_text(values, headers, row_number),
                                                {"table": table_number, "row": row_number, "headers": headers}, heading))
        if doc.inline_shapes or doc.element.xpath(".//w:drawing"):
            warnings.append("文档含图片，图片内文字未执行 OCR")
        return ParsedDocument(tuple(sections), tuple(warnings))

    @staticmethod
    def _row_text(values, headers, row_number):
        def value_text(value):
            return "" if value is None else str(value).strip()
        if row_number == 1:
            return " | ".join(value_text(value) for value in values)
        return " | ".join(f"{headers[index] if index < len(headers) and headers[index] else f'列{index + 1}'}: {value_text(value)}"
                          for index, value in enumerate(values))

    def _xlsx(self, data: bytes) -> ParsedDocument:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        sections, warnings = [], []
        try:
            for sheet in workbook.worksheets:
                if sheet.max_row * sheet.max_column > 2_000_000:
                    raise UnsupportedDocument("工作表单元格数量超过解析限制")
                headers = []
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    values = ["" if value is None else str(value) for value in row]
                    if number == 1:
                        headers = values
                    if any(values):
                        sections.append(Section(self._row_text(values, headers, number),
                                                {"sheet": sheet.title, "row": number, "headers": headers}, sheet.title))
        finally:
            workbook.close()
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if any(re.search(rb"<f(?:\s|>)", archive.read(name)) for name in archive.namelist()
                   if name.startswith("xl/worksheets/") and name.endswith(".xml")):
                warnings.append("公式仅使用文件内缓存值，不执行或重算；未缓存的结果可能为空")
            if any(name.startswith("xl/media/") for name in archive.namelist()):
                warnings.append("工作表含图片，未执行 OCR")
        return ParsedDocument(tuple(sections), tuple(warnings))

    def _pptx(self, data: bytes) -> ParsedDocument:
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(data))
        sections, warnings = [], []
        for number, slide in enumerate(presentation.slides, 1):
            heading = slide.shapes.title.text if slide.shapes.title else ""
            def visit(shapes):
                for shape in shapes:
                    if hasattr(shape, "shapes"):
                        visit(shape.shapes)
                    elif shape.has_table:
                        headers = [cell.text for cell in shape.table.rows[0].cells]
                        for row_number, row in enumerate(shape.table.rows, 1):
                            values = [cell.text for cell in row.cells]
                            if any(values):
                                sections.append(Section(self._row_text(values, headers, row_number),
                                                        {"slide": number, "shape": shape.shape_id, "row": row_number,
                                                         "headers": headers}, heading))
                    elif shape.has_text_frame and shape.text.strip():
                        sections.append(Section(shape.text.strip(), {"slide": number, "shape": shape.shape_id}, heading))
                    elif shape.shape_type == 13:
                        warnings.append(f"幻灯片 {number} 含图片，未执行 OCR")
                    elif getattr(shape, "has_chart", False):
                        warnings.append(f"幻灯片 {number} 含图表，未提取图表数据")
            visit(slide.shapes)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    sections.append(Section(notes, {"slide": number, "notes": True}, heading))
        return ParsedDocument(tuple(sections), tuple(dict.fromkeys(warnings)))

    def _pdf(self, data: bytes) -> ParsedDocument:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise UnsupportedDocument("不支持加密 PDF")
        sections, warnings = [], []
        for number, page in enumerate(reader.pages, 1):
            text = (page.extract_text() or "").strip()
            if text:
                sections.append(Section(text, {"page": number}))
                resources = page.get("/Resources", {})
                if hasattr(resources, "get_object"):
                    resources = resources.get_object()
                objects = resources.get("/XObject", {})
                if hasattr(objects, "get_object"):
                    objects = objects.get_object()
                if any(obj.get_object().get("/Subtype") == "/Image" for obj in objects.values()):
                    warnings.append(f"第 {number} 页含图片，图片内容未执行 OCR")
            else:
                warnings.append(f"第 {number} 页无可提取文字，未执行 OCR")
        return ParsedDocument(tuple(sections), tuple(warnings))

    @staticmethod
    def _rich_text(block):
        for value in block.values():
            if isinstance(value, dict) and isinstance(value.get("elements"), list):
                return "".join(str(item.get("text_run", {}).get("content", ""))
                               for item in value["elements"] if isinstance(item, dict)).strip()
        return ""

    def _native(self, payload: dict) -> ParsedDocument:
        if payload.get("kind") == "sheet":
            sections = []
            for sheet in payload["sheets"]:
                headers = sheet.get("headers", [])
                for record in sheet["rows"]:
                    number, values = record["row"], record["values"]
                    if any(value not in (None, "") for value in values):
                        sections.append(Section(self._row_text(values, headers, number),
                                                {"sheet": sheet["title"], "sheet_id": sheet["id"],
                                                 "row": number, "headers": headers}, sheet["title"]))
            return ParsedDocument(tuple(sections), tuple(payload.get("warnings", [])))
        if payload.get("kind") != "docx":
            raise UnsupportedDocument("不支持的飞书原生类型")
        blocks = {block["block_id"]: block for block in payload["blocks"]}
        if len(blocks) != len(payload["blocks"]):
            raise UnsupportedDocument("原生文档包含重复块编号")
        if any(child not in blocks for block in blocks.values() for child in block.get("children", [])):
            raise UnsupportedDocument("原生文档缺少子块，拒绝发布不完整正文")
        children = {child for block in blocks.values() for child in block.get("children", [])}
        roots = [key for key in blocks if key not in children]
        sections, warnings, visited = [], list(payload.get("warnings", [])), set()
        heading = ""
        def all_text(block_id, chain=frozenset()):
            if block_id in chain or block_id not in blocks:
                return ""
            block = blocks[block_id]
            return "\n".join(filter(None, [self._rich_text(block),
                *(all_text(child, chain | {block_id}) for child in block.get("children", []))]))
        def walk(block_id):
            nonlocal heading
            if block_id in visited:
                return
            visited.add(block_id)
            block = blocks[block_id]
            if block.get("block_type") == 31:
                table = block.get("table", {})
                cells = table.get("cells", block.get("children", []))
                columns = table.get("property", {}).get("column_size", 0)
                if not columns:
                    warnings.append(f"表格 {block_id} 缺少列信息")
                    return
                rows = [[all_text(cell) for cell in cells[start:start + columns]] for start in range(0, len(cells), columns)]
                headers = rows[0] if rows else []
                for row_number, values in enumerate(rows, 1):
                    if any(values):
                        sections.append(Section(self._row_text(values, headers, row_number),
                                                {"block_id": block_id, "row": row_number, "headers": headers}, heading))
                # Cells remain owned by this table, even if returned as top-level list entries.
                visited.update(cells)
                return
            text = self._rich_text(block)
            if 3 <= block.get("block_type", 0) <= 11:
                heading = text or heading
            if text:
                sections.append(Section(text, {"block_id": block_id}, heading))
            if block.get("block_type") in {23, 27}:
                warnings.append(f"块 {block_id} 为附件或图片，未执行 OCR/递归附件读取")
            for child in block.get("children", []):
                if child in blocks:
                    walk(child)
                else:
                    warnings.append(f"缺少子块 {child}")
        for root in roots:
            walk(root)
        return ParsedDocument(tuple(sections), tuple(dict.fromkeys(warnings)))
