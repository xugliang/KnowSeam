"""Actual vendored parser through Unix socket; no mocks for extraction."""
import io
import os
from pathlib import Path

import pytest

from knowseam.document_parsing.client import ParserClient
from knowseam.document_parsing.protocol import DoclingOptions
from knowseam.domain import SourceContent, UnsupportedDocument

pytestmark = pytest.mark.docling_real


@pytest.fixture
def client():
    path = os.environ.get('KNOWSEAM_TEST_PARSER_SOCKET')
    if not path:
        pytest.skip('true parser tests require isolated running socket and preloaded models')
    value = ParserClient(Path(path))
    value.status()
    try:
        yield value
    finally:
        value.close()


def test_docx_heading_and_table(client):
    from docx import Document
    doc = Document()
    doc.add_heading('示例功能', 1)
    doc.add_paragraph('示例产品支持企业文档检索')
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [('产品','数量'), ('示例产品','5')], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    stream = io.BytesIO()
    doc.save(stream)
    parsed = client.parse(SourceContent(stream.getvalue(), '说明.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'), DoclingOptions())
    assert '企业文档检索' in str(parsed.sections)
    assert any('示例产品' in section.text and '5' in section.text and section.locator.get('row') for section in parsed.sections)
    assert any(section.heading == '示例功能' for section in parsed.sections)


def test_xlsx_merged_header_and_pptx(client):
    from openpyxl import Workbook
    from pptx import Presentation
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '功能表'
    sheet['D10'] = '产品'
    sheet.merge_cells('D10:D11')
    sheet['E10'] = '财务'
    sheet.merge_cells('E10:F10')
    sheet['E11'], sheet['F11'] = '收入', '成本'
    sheet['D12'], sheet['E12'], sheet['F12'] = '示例产品', '5', '3'
    stream = io.BytesIO()
    workbook.save(stream)
    parsed = client.parse(SourceContent(stream.getvalue(), '功能.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'), DoclingOptions())
    row = next(s for s in parsed.sections if '示例产品' in s.text)
    assert row.text == '产品: 示例产品 | 财务 / 收入: 5 | 财务 / 成本: 3'
    assert row.locator['row'] == 12 and row.locator['column_start'] == 4
    assert row.locator['sheet'] == '功能表'
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = '示例产品'
    slide.placeholders[1].text = '支持企业知识问答'
    stream = io.BytesIO()
    presentation.save(stream)
    parsed = client.parse(SourceContent(stream.getvalue(), '演示.pptx', 'application/vnd.openxmlformats-officedocument.presentationml.presentation'), DoclingOptions())
    assert any('支持企业知识问答' in s.text for s in parsed.sections)
    assert any(s.locator.get('slide') == 1 for s in parsed.sections)


def chinese_image():
    from PIL import Image, ImageDraw, ImageFont
    font_path = os.environ.get('KNOWSEAM_TEST_CHINESE_FONT')
    if not font_path:
        pytest.skip('synthetic Chinese scan requires explicit CJK font file')
    image = Image.new('RGB', (1400, 500), 'white')
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(font_path, 48)
    draw.text((70,80), '示例企业知识库', font=font, fill='black')
    draw.text((70,190), '产品功能：支持文档检索', font=font, fill='black')
    return image


@pytest.mark.parametrize('format,media', [('PNG','image/png'), ('JPEG','image/jpeg'), ('TIFF','image/tiff'), ('PDF','application/pdf')])
def test_actual_chinese_scan(client, format, media):
    stream = io.BytesIO()
    chinese_image().save(stream, format=format)
    parsed = client.parse(SourceContent(stream.getvalue(), '扫描.'+format.lower(), media), DoclingOptions())
    text = '\n'.join(section.text for section in parsed.sections)
    assert '示例企业' in text and '文档检索' in text
    assert any(s.locator.get('page') == 1 for s in parsed.sections)


def test_empty_image_not_reported_as_success(client):
    from PIL import Image
    stream = io.BytesIO()
    Image.new('RGB', (600,400),'white').save(stream, format='PNG')
    with pytest.raises(UnsupportedDocument):
        client.parse(SourceContent(stream.getvalue(),'空白.png','image/png'),DoclingOptions())
