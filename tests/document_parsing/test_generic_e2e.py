"""Real PostgreSQL and source I/O; deterministic model inference, no company materials."""
from dataclasses import replace
from pathlib import Path
import os

import httpx
import pytest

from test_postgres_e2e import case as _case, checked_dsn as _checked_dsn, tokenizer as _tokenizer
from knowseam.domain import UnsupportedDocument
from knowseam.document_parsing.config import DoclingOptions
from knowseam.document_parsing.docling_adapter import DoclingAdapter
from knowseam.document_parsing.protocol import ParserIdentity
from knowseam.parsers import StructuredDocumentParser

pytestmark = pytest.mark.postgres
case, checked_dsn, tokenizer = _case, _checked_dsn, _tokenizer


class Client:
    identity = ParserIdentity(upstream_sha='a'*40, runtime_source_sha256='b'*64,
                              dependency_fingerprint='c'*64, model_fingerprint='d'*64)
    fail = False
    callback = None
    def status(self):
        return self.identity
    def parse(self, content, options):
        if self.callback:
            self.callback()
        if self.fail:
            raise UnsupportedDocument('parser_timeout')
        result = StructuredDocumentParser().parse(content)
        from knowseam.document_parsing.fingerprint import parser_fingerprint
        from knowseam.document_parsing.normalizer import NORMALIZER_VERSION
        return replace(result, parser_version=parser_fingerprint(self.identity, options,
                       normalizer_version=NORMALIZER_VERSION))


def test_docling_pipeline_change_reindexes_without_reembedding(case):
    client = Client()
    case.worker.parser = DoclingAdapter(client, DoclingOptions())
    assert case.sync()['published'] == 1
    before, calls = case.document(), len(case.embedder.calls)
    assert case.sync()['unchanged'] == 1
    client.identity = client.identity.model_copy(update={'runtime_source_sha256': 'e'*64})
    case.worker.parser = DoclingAdapter(client, DoclingOptions())
    assert case.sync()['published'] == 1
    assert case.document()['active_version'] != before['active_version']
    assert len(case.embedder.calls) == calls


def test_docling_failure_keeps_old_version_stale(case):
    assert case.sync()['published'] == 1
    before = case.document()
    client = Client()
    client.fail = True
    case.worker.parser = DoclingAdapter(client, DoclingOptions())
    assert case.sync()['failed'] == 1
    after = case.document()
    assert after['active_version'] == before['active_version'] and after['status'] == 'stale'


def test_first_failure_cannot_be_read(case):
    client = Client()
    client.fail = True
    case.worker.parser = DoclingAdapter(client, DoclingOptions())
    assert case.sync()['failed'] == 1
    assert not case.document()['active_version'] and not case.search()['hits']


def test_rollback_uses_current_source_revision(case):
    case.worker.parser = DoclingAdapter(Client(), DoclingOptions())
    assert case.sync()['published'] == 1
    case.write('当前有效内容 ROLLBACK-CURRENT，不可恢复旧内容。')
    case.worker.parser = StructuredDocumentParser()
    assert case.sync()['published'] == 1
    assert 'ROLLBACK-CURRENT' in str(case.read())
    assert 'TEST-00000001' not in str(case.read())


@pytest.mark.parametrize('change', ['revoke', 'delete'])
def test_revoke_or_delete_during_parse_cannot_publish(case, change):
    assert case.sync()['published'] == 1
    client = Client()
    doc = case.document()
    def revoke():
        with case.store.connect() as connection:
            connection.execute('UPDATE sk_documents SET public=false, status=%s WHERE id=%s',
                               ('deleted' if change == 'delete' else 'unavailable', doc['id']))
    client.callback = revoke
    case.worker.parser = DoclingAdapter(client, DoclingOptions())
    assert case.sync()['failed'] == 1
    assert not case.search()['hits']
    assert case.document()['active_version'] == doc['active_version']


@pytest.mark.asyncio
async def test_real_source_index_rest_and_mcp(case):
    assert await case.worker.run(case.source, case.connector, config_hash=case.config_hash) == {
        'seen': 1, 'published': 1, 'unchanged': 0, 'unavailable': 0, 'failed': 0, 'deleted': 0}
    await assert_public_api_and_mcp(case, 'TEST-00000001')


async def assert_public_api_and_mcp(case, query):
    from knowseam.api import create_app
    from knowseam.config import ServiceClient
    from knowseam.mcp_server import GatewayConfig, create_app as gateway
    from test_mcp_gateway import http_client, call
    import hashlib
    token = 'generic-e2e-readonly-token-ABCDEFGHIJKLMNOPQRSTUVWXYZ-1234567890'
    client = ServiceClient(token, case.principal.product_id, case.company, (case.collection,), (case.principal.actor_id,))
    api = create_app(case.service, (client,))
    headers = {'Authorization': 'Bearer '+token, 'X-Actor-Id': case.principal.actor_id,
               'X-Company-Id': case.company, 'X-Product-Id': case.principal.product_id}
    async with httpx.AsyncClient(transport=httpx.ASGITransport(api), base_url='http://knowledge', headers=headers) as http:
        searched = await http.post('/v1/search', json={'query': query, 'collection_ids': [case.collection]})
        assert searched.status_code == 200 and searched.json()['hits']
        hit = searched.json()['hits'][0]
        read = await http.get('/v1/documents/'+hit['document_id'], params={'version_id': hit['version_id']})
        assert read.status_code == 200 and read.json()['sections']
        assert case.reranker.calls
    config = GatewayConfig.model_validate({'upstream_url': 'http://knowledge', 'upstream_token': token,
        'allow_insecure_http': True, 'allowed_hosts': ['testserver'], 'bindings': [{
            'token_sha256': hashlib.sha256(token.encode()).hexdigest(), 'actor_id': case.principal.actor_id,
            'company_id': case.company, 'product_id': case.principal.product_id,
            'allowed_collections': [case.collection]}]})
    app = gateway(config, transport=httpx.ASGITransport(api))
    async with http_client(app, token=token) as http:
        result = await call(http, 'knowledge_search', {'query': query, 'collection_ids': [case.collection]})
        assert result['structuredContent']['hits']


@pytest.mark.docling_real
@pytest.mark.asyncio
async def test_real_docling_office_source_to_database(case):
    from knowseam.document_parsing.client import ParserClient
    socket = os.environ.get('KNOWSEAM_TEST_PARSER_SOCKET')
    if not socket:
        pytest.skip('requires isolated running parser socket')
    import io
    from docx import Document
    doc = Document()
    doc.add_heading('示例企业功能', 1)
    doc.add_paragraph('示例产品支持 REAL-PARSER-INDEX 企业检索。')
    stream = io.BytesIO()
    doc.save(stream)
    (case.root/'guide.docx').write_bytes(stream.getvalue())
    (case.root/'guide.md').unlink()
    case.connector.public_files |= {'guide.docx'}
    client = ParserClient(Path(socket))
    try:
        case.worker.parser = DoclingAdapter(client, DoclingOptions())
        counts = await case.worker.run(case.source, case.connector, config_hash=case.config_hash)
        assert counts['published'] == 1
        await assert_public_api_and_mcp(case, 'REAL-PARSER-INDEX')
    finally:
        client.close()
