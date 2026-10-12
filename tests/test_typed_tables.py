import io
import json
import zipfile
from dataclasses import replace

import pytest

from knowseam.domain import SourceContent, UnsupportedDocument
from knowseam.parsers import StructuredDocumentParser


def workbook_bytes():
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "价格"
    sheet.append(["编号", "金额（元）", "数值", "公式"])
    sheet.append(["0012", 1, 0, "=B2+C2"])
    sheet.merge_cells("A3:B4")
    sheet["A3"] = "合并值"
    out = io.BytesIO()
    book.save(out)
    original, changed = zipfile.ZipFile(io.BytesIO(out.getvalue())), io.BytesIO()
    with zipfile.ZipFile(changed, "w") as archive:
        for name in original.namelist():
            data = original.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<v>1</v>", b"<v>-1234567890123456789.00100</v>")
            archive.writestr(name, data)
    original.close()
    return changed.getvalue()


def test_native_xlsx_exact_decimal_leading_zero_formula_and_merge():
    parser = StructuredDocumentParser(structured_tables=True)
    doc = parser.parse(SourceContent(workbook_bytes(), "prices.xlsx", "application/octet-stream"))
    assert doc.structure_version == 2 and parser.version != "structured-v1"
    table = doc.tables[0]
    cells = {(c.row, c.column): c for c in table.cells}
    assert cells[2, 1].value == "0012" and cells[2, 1].value_kind == "text"
    assert cells[2, 2].value == "-1234567890123456789.00100"
    assert cells[2, 2].unit == "元"
    assert cells[2, 3].value == "0" and cells[2, 3].value_kind == "number"
    assert cells[2, 4].formula == "=B2+C2" and cells[2, 4].formula_state == "cache_missing"
    assert (cells[3, 1].row_span, cells[3, 1].column_span) == (2, 2)
    assert (4, 2) not in cells


def test_protocol_v1_unchanged_v2_roundtrip_and_hash():
    from knowseam.document_parsing.protocol import encode_document, decode_document
    from knowseam.chunking import parsed_hash

    doc = StructuredDocumentParser(structured_tables=True).parse(
        SourceContent(workbook_bytes(), "p.xlsx", "application/octet-stream")
    )
    v1 = encode_document(doc)
    assert set(json.loads(v1)) == {"sections", "warnings", "parser_version"}
    assert decode_document(v1).tables == () and decode_document(v1).structure_version == 1
    v2 = encode_document(doc, protocol_version=2)
    assert decode_document(v2, protocol_version=2) == doc
    with pytest.raises(UnsupportedDocument):
        decode_document(v2)
    changed = doc.tables[0].model_copy(
        update={
            "cells": tuple(
                c.model_copy(update={"unit": "万元"}) if c.unit else c for c in doc.tables[0].cells
            )
        }
    )
    assert parsed_hash(replace(doc, tables=(changed,))) != parsed_hash(doc)


def test_types_reject_float_coordinates_and_ambiguous_empty():
    from knowseam.tables.domain import Cell
    from pydantic import ValidationError

    for changes in (
        {"row": True},
        {"row": 0},
        {"value_kind": "number", "value": 1.1},
        {"value_kind": "empty", "value": "0"},
        {"value_kind": "number", "value": "NaN"},
    ):
        with pytest.raises(ValidationError):
            Cell(
                row=changes.get("row", 1),
                column=1,
                raw_text="",
                **{k: v for k, v in changes.items() if k != "row"},
            )


def test_xlsx_styles_do_not_scan_for_each_cell_and_chartsheets_are_skipped(monkeypatch):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    from knowseam.tables.native import xlsx_tables

    original, opened = ReadOnlyWorksheet._get_source, []

    def counted(self):
        opened.append(self.title)
        return original(self)

    monkeypatch.setattr(ReadOnlyWorksheet, "_get_source", counted)
    book = Workbook()
    sheet = book.active
    for n in range(200):
        sheet.append([n, n + 1])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=1, min_row=1, max_row=10))
    book.create_chartsheet("chart", 0).add_chart(chart)
    data = io.BytesIO()
    book.save(data)
    tables = xlsx_tables(data.getvalue())
    assert len(tables) == 1 and len(tables[0].cells) == 400
    assert len(opened) <= 2
