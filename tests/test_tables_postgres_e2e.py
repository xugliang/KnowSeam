import pytest

from test_postgres_e2e import case as _case, checked_dsn as _checked_dsn, tokenizer as _tokenizer

from test_typed_tables import workbook_bytes
from knowseam.migrations.runner import MigrationRunner
from knowseam.domain import NotFound

case, checked_dsn, tokenizer = _case, _checked_dsn, _tokenizer


def setup_tables(case):
    from knowseam.config import FeaturesConfig
    from knowseam.publication import build_publisher
    from knowseam.parsers import StructuredDocumentParser

    MigrationRunner(case.store.connect).apply()
    case.worker.structure_publisher = build_publisher(case.store, FeaturesConfig(tables=True))
    case.worker.parser = StructuredDocumentParser(structured_tables=True)
    (case.root / "prices.xlsx").write_bytes(workbook_bytes())
    case.connector.public_files = frozenset(("guide.md", "prices.xlsx"))
    assert case.sync()["published"] == 2
    with case.store.connect() as c:
        return c.execute(
            "SELECT * FROM sk_documents WHERE source_id=%s AND external_id=%s", (case.source, "prices.xlsx")
        ).fetchone()


def test_table_only_publication_and_revocation(case, monkeypatch):
    from knowseam.tables.contracts import TableListRequest, TableReadRequest
    from knowseam.tables.postgres import PostgresTableReader

    doc = setup_tables(case)
    reader = PostgresTableReader(case.store)
    identity = dict(document_id=doc["id"], version_id=doc["active_version"])
    assert reader.list_tables(case.principal, TableListRequest(**identity))["tables"][0]["table_id"] == "t-1"
    request = TableReadRequest(**identity, table_id="t-1", row_start=2, row_limit=1)
    result = reader.read_table(case.principal, request)
    assert next(c for c in result["cells"] if c["column"] == 2)["value"] == "-1234567890123456789.00100"
    import knowseam.tables.postgres as module

    original = module.read_table

    def revoke(*args):
        result = original(*args)
        case.store.fail_document(doc["id"], "test", revoke=True)
        return result

    monkeypatch.setattr(module, "read_table", revoke)
    with pytest.raises(NotFound):
        reader.read_table(case.principal, request)


def test_table_failure_rolls_back_active_version(case, monkeypatch):
    doc = setup_tables(case)
    import knowseam.tables.postgres as module

    def fail(*args):
        raise ValueError("injected table failure")

    monkeypatch.setattr(module, "write_tables", fail)
    # A forced scan with the same file is unchanged; mutate the title-bearing source bytes.
    before = doc["active_version"]
    from openpyxl import load_workbook

    book = load_workbook(case.root / "prices.xlsx")
    book.active["B2"] = 10
    book.save(case.root / "prices.xlsx")
    assert case.sync()["failed"] == 1
    with case.store.connect() as c:
        current = c.execute("SELECT active_version FROM sk_documents WHERE id=%s", (doc["id"],)).fetchone()
    assert current["active_version"] == before
