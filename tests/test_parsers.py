import io
import json
import unittest
import zipfile

from knowseam.domain import SourceContent, UnsupportedDocument
from knowseam.parsers import FEISHU_NATIVE_TYPE, StructuredDocumentParser


class ParserTests(unittest.TestCase):
    def setUp(self):
        self.parser = StructuredDocumentParser()

    def parse(self, data, name, media="application/octet-stream"):
        return self.parser.parse(SourceContent(data, name, media))

    def test_markdown_keeps_heading_and_original_line_numbers(self):
        result = self.parse("# 产品说明\n\n产品可导出报告。\n第二行。\n\n![架构](image.png)".encode(), "说明.md")
        self.assertEqual(result.sections[1].heading, "产品说明")
        self.assertEqual(result.sections[1].locator, {"line_start": 3, "line_end": 4})
        self.assertTrue(result.warnings)

    def test_docx_preserves_paragraph_and_table_header(self):
        from docx import Document
        document = Document()
        document.add_heading("公司知识", 1)
        document.add_paragraph("产品支持导出。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "产品", "功能"
        table.cell(1, 0).text, table.cell(1, 1).text = "Kino", "视频生成"
        stream = io.BytesIO()
        document.save(stream)
        result = self.parse(stream.getvalue(), "产品.docx")
        row = next(section for section in result.sections if section.locator.get("row") == 2)
        self.assertIn("产品: Kino", row.text)
        self.assertEqual(row.heading, "公司知识")
        self.assertEqual(row.locator["table"], 1)

    def test_xlsx_multiple_sheets_rows_and_formula_not_executed(self):
        from openpyxl import Workbook
        workbook = Workbook()
        first = workbook.active
        first.title = "Kino产品线"
        first.append(["产品", "价格"])
        first.append(["Kino", "=123+456"])
        second = workbook.create_sheet("示例产品产品线")
        second.append(["功能", "介绍"])
        second.append(["标书", "自动整理"])
        stream = io.BytesIO()
        workbook.save(stream)
        result = self.parse(stream.getvalue(), "产品.xlsx")
        self.assertEqual({s.locator["sheet"] for s in result.sections}, {"Kino产品线", "示例产品产品线"})
        self.assertTrue(any("缓存值" in warning for warning in result.warnings))
        self.assertNotIn("579", "\n".join(section.text for section in result.sections))

    def test_pptx_slide_table_and_notes(self):
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "产品能力"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(5), Inches(2)).table
        table.cell(0, 0).text, table.cell(0, 1).text = "模块", "能力"
        table.cell(1, 0).text, table.cell(1, 1).text = "标书", "检索"
        slide.notes_slide.notes_text_frame.text = "销售备注"
        stream = io.BytesIO()
        deck.save(stream)
        result = self.parse(stream.getvalue(), "产品.pptx")
        self.assertTrue(all(section.locator["slide"] == 1 for section in result.sections))
        self.assertTrue(any(section.locator.get("notes") and "销售备注" in section.text for section in result.sections))
        self.assertTrue(any("模块: 标书" in section.text for section in result.sections))

    def test_pdf_text_has_page_locator(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
            DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 20 150 Td (Product knowledge) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        stream = io.BytesIO()
        writer.write(stream)
        result = self.parse(stream.getvalue(), "text.pdf")
        self.assertEqual(result.sections[0].locator, {"page": 1})
        self.assertIn("Product knowledge", result.sections[0].text)

    def test_blank_scan_pdf_and_images_are_not_claimed_ocr(self):
        from pypdf import PdfWriter
        writer, stream = PdfWriter(), io.BytesIO()
        writer.add_blank_page(width=200, height=200)
        writer.write(stream)
        with self.assertRaisesRegex(UnsupportedDocument, "OCR"):
            self.parse(stream.getvalue(), "scan.pdf")
        with self.assertRaises(UnsupportedDocument):
            self.parse(b"fake image", "image.png")

    def test_macro_and_zip_expansion_limit_rejected(self):
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as stream:
            stream.writestr("word/vbaProject.bin", b"macro")
        with self.assertRaisesRegex(UnsupportedDocument, "宏"):
            self.parse(archive.getvalue(), "unsafe.docx")
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as stream:
            stream.writestr("large.xml", b"a" * 1000)
        parser = StructuredDocumentParser(max_expanded_bytes=100)
        with self.assertRaisesRegex(UnsupportedDocument, "展开"):
            parser.parse(SourceContent(archive.getvalue(), "bomb.xlsx", ""))

    def test_native_docx_block_tree_table_and_image_warning(self):
        def text(block_id, value):
            return {"block_id": block_id, "block_type": 2, "text": {"elements": [{"text_run": {"content": value}}]}}
        blocks = [{"block_id": "root", "block_type": 1, "children": ["h", "t", "image"]},
                  {"block_id": "h", "block_type": 3, "heading1": {"elements": [{"text_run": {"content": "产品"}}]}},
                  {"block_id": "t", "block_type": 31, "children": ["a", "b", "c", "d"],
                   "table": {"cells": ["a", "b", "c", "d"], "property": {"column_size": 2}}},
                  *[{"block_id": cell, "block_type": 32, "children": [cell + "1"]} for cell in "abcd"],
                  text("a1", "名称"), text("b1", "功能"), text("c1", "Kino"), text("d1", "视频"),
                  {"block_id": "image", "block_type": 27, "image": {"token": "secret"}}]
        result = self.parse(json.dumps({"kind": "docx", "blocks": blocks}).encode(), "native.json", FEISHU_NATIVE_TYPE)
        self.assertEqual(len(result.sections), 3)
        self.assertEqual(result.sections[-1].locator["row"], 2)
        self.assertIn("名称: Kino", result.sections[-1].text)
        self.assertTrue(any("OCR" in warning for warning in result.warnings))

    def test_native_sheet_keeps_sheet_row_and_headers(self):
        payload = {"kind": "sheet", "sheets": [{"id": "abc", "title": "Kino", "headers": ["模块", "描述"],
                    "rows": [{"row": 5, "values": ["视频", "内容生成"]}]}]}
        result = self.parse(json.dumps(payload).encode(), "native.json", FEISHU_NATIVE_TYPE)
        self.assertEqual(result.sections[0].locator["row"], 5)
        self.assertIn("模块: 视频", result.sections[0].text)

    def test_unsupported_legacy_and_corrupt_input(self):
        for filename, data in [("legacy.ppt", b"old"), ("bad.docx", b"not zip"), ("bad.txt", b"\0a")]:
            with self.subTest(filename=filename), self.assertRaises(UnsupportedDocument):
                self.parse(data, filename)


if __name__ == "__main__":
    unittest.main()
